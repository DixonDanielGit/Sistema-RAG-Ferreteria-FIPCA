"""Pruebas unitarias de extraccion, normalizacion y persistencia ETL."""

from __future__ import annotations

import hashlib
import io
import zipfile
from decimal import Decimal

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.config import get_settings
from app.models import inventario as inventario_model
from app.services import etl_service


def _bytes_excel(filas):
    libro = Workbook()
    hoja = libro.active
    for fila in filas:
        hoja.append(fila)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _bytes_word(filas):
    documento = Document()
    tabla = documento.add_table(rows=0, cols=len(filas[0]))
    for fila in filas:
        celdas = tabla.add_row().cells
        for celda, valor in zip(celdas, fila):
            celda.text = str(valor)
    salida = io.BytesIO()
    documento.save(salida)
    return salida.getvalue()


def test_txt_normaliza_cabecera_y_precio_venezolano():
    contenido = (
        "Código;Nombre;Precio;Marca;Descripción\n"
        "SKU-01;Martillo;1.234,50;FIPCA;Acero\n"
    ).encode()

    productos, errores, detalles = etl_service._extraer_txt(contenido)

    assert errores == 0
    assert detalles == []
    assert productos == [
        {
            "codigo": "SKU-01",
            "nombre": "Martillo",
            "marca": "FIPCA",
            "precio": Decimal("1234.50"),
            "descuentos": Decimal("0.00"),
            "descripcion": "Acero",
        }
    ]


def test_txt_acepta_orden_posicional_documentado():
    productos, errores, _ = etl_service._extraer_txt(
        b"SKU-01;Martillo;12,50;Acero;FIPCA\n"
    )

    assert errores == 0
    assert productos[0]["codigo"] == "SKU-01"
    assert productos[0]["descripcion"] == "Acero"
    assert productos[0]["marca"] == "FIPCA"
    assert productos[0]["precio"] == Decimal("12.50")


def test_txt_posicional_con_comas_ambiguas_se_rechaza():
    with pytest.raises(etl_service.ErrorIngesta, match="punto y coma"):
        etl_service._extraer_txt(b"SKU-1,Martillo,1,25\n")


def test_txt_informa_filas_invalidas_sin_perder_las_validas():
    contenido = b"codigo;nombre;precio\nA1;Martillo;12,50\nA2;Clavo;-1\n"

    productos, errores, detalles = etl_service._extraer_txt(contenido)

    assert len(productos) == 1
    assert errores == 1
    assert "Fila 3" in detalles[0]


def test_normaliza_codigo_y_espacios_de_producto():
    productos, _, _ = etl_service._extraer_txt(
        b"codigo;nombre;precio;descripcion\nsku-01;  Llave   ajustable ;2,50; acero   inoxidable \n"
    )

    assert productos[0]["codigo"] == "SKU-01"
    assert productos[0]["nombre"] == "Llave ajustable"
    assert productos[0]["descripcion"] == "acero inoxidable"


def test_archivo_valido_guarda_original_y_resultado_con_uuid(monkeypatch, tmp_path):
    monkeypatch.setenv("UPLOADS_DIR", str(tmp_path / "cargas"))
    get_settings.cache_clear()
    contenido = b"codigo;nombre;precio\nSKU-31;Pinza;3.25\n"
    creado = {
        "id_archivo": 31,
        "nombre": "catalogo.txt",
        "tipo": "txt",
        "tamano_bytes": len(contenido),
    }
    estados = []
    productos_insertados = []
    datos_archivo = {}

    def crear(**valores):
        datos_archivo.update(valores)
        return creado

    def actualizar(id_archivo, **campos):
        estados.append((id_archivo, campos))
        return {
            **creado,
            "estado": campos["estado"],
            "filas_ok": campos.get("filas_ok", 0),
            "filas_error": campos.get("filas_error", 0),
            "detalle": campos.get("detalle"),
        }

    monkeypatch.setattr(inventario_model, "crear_archivo", crear)
    monkeypatch.setattr(inventario_model, "actualizar_estado_archivo", actualizar)
    monkeypatch.setattr(
        inventario_model,
        "insertar_o_actualizar_productos",
        lambda id_archivo, productos: productos_insertados.extend(productos) or len(productos),
    )
    try:
        resultado = etl_service.procesar_archivo(
            r"C:\proveedores\catalogo.txt", contenido, 7
        )
    finally:
        get_settings.cache_clear()

    directorio = tmp_path / "cargas"
    original = list(directorio.glob("*.txt"))
    assert resultado["estado"] == "ok"
    assert resultado["filas_ok"] == 1
    assert len(original) == 1
    assert original[0].name != "catalogo.txt"
    assert original[0].read_bytes() == contenido
    assert datos_archivo["ruta"] == original[0].name
    assert datos_archivo["hash_sha256"] == hashlib.sha256(contenido).hexdigest()
    assert productos_insertados[0]["codigo"] == "SKU-31"
    assert estados[0][1]["estado"] == "procesando"


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        ("$ 1,234.56", Decimal("1234.56")),
        ("Bs. 1.234,56", Decimal("1234.56")),
        ("0", Decimal("0.00")),
    ],
)
def test_normaliza_formatos_de_precio(valor, esperado):
    assert etl_service._leer_precio(valor) == esperado


@pytest.mark.parametrize("valor", ["", "-1", "1.2.3", "(1,00)", "NaN"])
def test_rechaza_precios_invalidos_o_negativos(valor):
    with pytest.raises(etl_service.ErrorIngesta):
        etl_service._leer_precio(valor)


def test_excel_extrae_tabla_con_cabecera():
    contenido = _bytes_excel(
        [
            ["Código", "Producto", "Precio", "Marca"],
            ["SKU-02", "Taladro", 19.95, "FIPCA"],
        ]
    )

    productos, errores, _ = etl_service._extraer_xlsx(contenido)

    assert errores == 0
    assert productos[0]["codigo"] == "SKU-02"
    assert productos[0]["precio"] == Decimal("19.95")


def test_excel_omite_hoja_sin_catalogo_y_lee_la_hoja_siguiente():
    libro = Workbook()
    libro.active.title = "Instrucciones"
    libro.active.append(["Lista de precios"])
    hoja = libro.create_sheet("Catalogo")
    hoja.append(["SKU", "Nombre", "PVP"])
    hoja.append(["SKU-03", "Llave", 5])
    salida = io.BytesIO()
    libro.save(salida)

    productos, errores, _ = etl_service._extraer_xlsx(salida.getvalue())

    assert errores == 0
    assert [producto["codigo"] for producto in productos] == ["SKU-03"]


def test_docx_extrae_tabla_con_cabecera():
    contenido = _bytes_word(
        [["Referencia", "Artículo", "Precio"], ["SKU-04", "Destornillador", "7,25"]]
    )

    productos, errores, _ = etl_service._extraer_docx(contenido)

    assert errores == 0
    assert productos[0]["codigo"] == "SKU-04"
    assert productos[0]["precio"] == Decimal("7.25")


def test_pdf_sin_capa_de_texto_se_rechaza():
    escritor = PdfWriter()
    escritor.add_blank_page(width=200, height=200)
    salida = io.BytesIO()
    escritor.write(salida)

    with pytest.raises(etl_service.ErrorIngesta, match="no contiene texto"):
        etl_service._extraer_pdf(salida.getvalue())


def test_pdf_extrae_filas_con_cabecera_y_nombres_compuestos(monkeypatch):
    class Pagina:
        def extract_text(self):
            return "Código Producto Precio\nSKU-1 Martillo acero 2,50"

    class Lector:
        pages = [Pagina()]

    monkeypatch.setattr(etl_service, "PdfReader", lambda *_args, **_kwargs: Lector())

    productos, errores, _ = etl_service._extraer_pdf(b"pdf de prueba")

    assert errores == 0
    assert productos[0]["codigo"] == "SKU-1"
    assert productos[0]["nombre"] == "Martillo acero"
    assert productos[0]["precio"] == Decimal("2.50")


def test_rechaza_archivo_xlsx_con_zip_danado():
    with pytest.raises(etl_service.ErrorIngesta, match="comprimido"):
        etl_service._validar_paquete(b"no es un zip")


def test_rechaza_zip_bomb_por_razon_de_compresion():
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", compression=zipfile.ZIP_DEFLATED) as archivo:
        archivo.writestr("xl/workbook.xml", b"0" * (2 * 1024 * 1024))

    with pytest.raises(etl_service.ErrorIngesta, match="compresion seguro"):
        etl_service._validar_paquete(salida.getvalue())


def test_rechaza_mas_de_20000_filas():
    filas = [["codigo", "nombre", "precio"]]
    filas.extend((f"SKU-{i}", "Producto", "1.00") for i in range(20_001))

    with pytest.raises(etl_service.ErrorIngesta, match="20000 filas"):
        etl_service._extraer_filas(filas)


def test_txt_de_5000_productos_se_procesa_completo():
    filas = [["codigo", "nombre", "precio"]]
    filas.extend((f"SKU-{i}", f"Producto {i}", "1,25") for i in range(5_000))

    productos, errores, _ = etl_service._extraer_filas(filas)

    assert len(productos) == 5_000
    assert errores == 0


def test_formatos_no_admitidos_se_rechazan_sin_guardar_archivo():
    resultado = etl_service.procesar_archivo("catalogo.csv", b"a,b,c", 1)

    assert resultado["estado"] == "rechazado"
    assert resultado["id_archivo"] is None


def test_limite_de_20_mb_se_aplica_antes_de_guardar(monkeypatch):
    monkeypatch.setattr(etl_service, "MAX_ARCHIVO_BYTES", 3)

    resultado = etl_service.procesar_archivo("catalogo.txt", b"1234", 1)

    assert resultado["estado"] == "rechazado"
    assert "20 MB" in resultado["detalle"]


def test_upsert_no_actualiza_stock_y_invalida_vector_si_cambia_descripcion(monkeypatch):
    capturado = {}

    class CursorFalso:
        pass

    class ContextoCursor:
        def __enter__(self):
            return CursorFalso()

        def __exit__(self, *_):
            return False

    def ejecutar_valores(cursor, consulta, valores, page_size):
        capturado["consulta"] = consulta.lower()
        capturado["valores"] = valores

    monkeypatch.setattr(inventario_model, "cursor_rag", ContextoCursor)
    monkeypatch.setattr(inventario_model, "execute_values", ejecutar_valores)

    total = inventario_model.insertar_o_actualizar_productos(
        5,
        [
            {
                "codigo": "SKU-05",
                "nombre": "Llave",
                "precio": Decimal("2.50"),
            }
        ],
    )

    assert total == 1
    assert "stock" not in capturado["consulta"]
    assert "descripcion_vectorial = case" in capturado["consulta"]
    assert capturado["valores"][0][0] == 5
    assert capturado["valores"][0][5] == Decimal("16.00")


def test_migracion_etl_es_idempotente_y_sin_dependencia_fk_usuario():
    raiz = etl_service.Path(__file__).resolve().parents[2]
    sql = (raiz / "scripts" / "migracion_3_0.sql").read_text(encoding="utf-8").lower()

    assert "add column if not exists" in sql
    assert "create index if not exists idx_archivos_hash" in sql
    assert "id_usuario int" in sql
    assert "references usuarios" not in sql
