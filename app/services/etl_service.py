"""Extraccion, normalizacion y carga de catalogos de proveedores."""

from __future__ import annotations

import csv
import hashlib
import io
import itertools
import logging
import os
import re
import unicodedata
import uuid
import zipfile
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable
from xml.etree.ElementTree import ParseError

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.config import get_settings
from app.db.session import BaseDeDatosNoDisponible
from app.models import inventario as inventario_model
from psycopg2 import Error as ErrorPostgres

logger = logging.getLogger("ferreteria.etl")

EXTENSIONES_PERMITIDAS = frozenset({".pdf", ".xlsx", ".docx", ".txt"})
MAX_ARCHIVO_MB = 20
MAX_ARCHIVO_BYTES = MAX_ARCHIVO_MB * 1024 * 1024
MAX_ARCHIVOS_INGESTA = 20
MAX_FILAS_ARCHIVO = 20_000
MAX_PAGINAS_PDF = 300
MAX_ENTRADAS_ZIP = 500
MAX_DESCOMPRIMIDO_ZIP = 100 * 1024 * 1024
MAX_RAZON_COMPRESION = 100
MAX_DETALLE = 4000
MAX_ERRORES_DETALLE = 20

# La cabecera de PDF/XLSX/DOCX es obligatoria. TXT puede usar el orden indicado
# en la documentacion cuando no contiene una cabecera reconocible.
TXT_COLUMNAS_POSICIONALES = ("codigo", "nombre", "precio", "descripcion", "marca")

SINONIMOS = {
    "codigo": {
        "codigo", "cod", "sku", "referencia", "ref", "item", "itemnumber",
        "productcode", "code",
    },
    "nombre": {"nombre", "producto", "articulo", "product", "name"},
    "precio": {
        "precio", "preciounitario", "precioventa", "pvp", "costo", "cost",
        "price", "unitprice",
    },
    "marca": {"marca", "fabricante", "brand", "manufacturer"},
    "descripcion": {"descripcion", "description", "detalle", "observacion", "notes"},
    "descuentos": {"descuento", "descuentos", "discount", "discounts"},
}
_POR_SINONIMO = {
    sinonimo: campo
    for campo, sinonimos in SINONIMOS.items()
    for sinonimo in sinonimos
}


class ErrorIngesta(ValueError):
    """Error de formato o contenido que debe aislarse a un archivo."""


class CabeceraIngestaAusente(ErrorIngesta):
    """La hoja o tabla no contiene una cabecera de productos reconocible."""


class ArchivoSinFilasValidas(ErrorIngesta):
    """Todas las filas del catalogo fallaron la validacion."""

    def __init__(self, cantidad: int, detalle: str) -> None:
        self.filas_error = cantidad
        super().__init__(detalle)


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _normalizar_texto(valor: Any) -> str:
    texto = _texto(valor)
    if any(ord(caracter) < 32 and caracter not in "\t\r\n" for caracter in texto):
        raise ErrorIngesta("El texto contiene caracteres de control")
    return " ".join(texto.split())


def _clave(valor: Any) -> str:
    texto = unicodedata.normalize("NFKD", _texto(valor).casefold())
    sin_tildes = "".join(caracter for caracter in texto if not unicodedata.combining(caracter))
    return re.sub(r"[^a-z0-9]+", "", sin_tildes)


def _mapear_cabecera(fila: Iterable[Any]) -> dict[str, int] | None:
    columnas: dict[str, int] = {}
    for indice, valor in enumerate(fila):
        campo = _POR_SINONIMO.get(_clave(valor))
        if campo:
            if campo in columnas:
                raise ErrorIngesta(f"La cabecera repite la columna '{campo}'")
            columnas[campo] = indice
    if {"codigo", "nombre", "precio"}.issubset(columnas):
        return columnas
    return None


def _separador(texto: str) -> str | None:
    for separador in ("\t", ";", "|"):
        if separador in texto:
            return separador
    return None


def _filas_delimitadas(texto: str) -> list[list[str]]:
    lineas: list[str] = []
    for numero, linea in enumerate(io.StringIO(texto), start=1):
        if numero > MAX_FILAS_ARCHIVO + 101:
            raise ErrorIngesta(f"El archivo supera el maximo de {MAX_FILAS_ARCHIVO} filas")
        if linea.strip():
            lineas.append(linea.rstrip("\r\n"))
    if not lineas:
        return []
    separador = _separador(lineas[0])
    if separador is None:
        separador = max(
            ("\t", ";", "|", ","),
            key=lambda candidato: sum(linea.count(candidato) for linea in lineas[:10]),
        )
        if not any(separador in linea for linea in lineas):
            return [[linea.strip()] for linea in lineas]
    return [next(csv.reader([linea], delimiter=separador, skipinitialspace=True)) for linea in lineas]


def _leer_precio(valor: Any, *, campo: str = "precio") -> Decimal:
    if isinstance(valor, bool):
        raise ErrorIngesta(f"{campo.capitalize()} invalido")
    if isinstance(valor, (int, float, Decimal)):
        texto = str(valor)
    else:
        texto = _texto(valor)
    if "(" in texto or ")" in texto:
        raise ErrorIngesta(f"{campo.capitalize()} no puede usar notacion negativa")
    texto = re.sub(r"(?i)(?:US\$|USD|VES|BS(?:\.F)?\.?)", "", texto)
    texto = texto.replace("$", "").replace(" ", "").replace("\u00a0", "")
    if re.search(r"[^0-9,.\-]", texto):
        raise ErrorIngesta(f"{campo.capitalize()} invalido")
    texto = re.sub(r"[^0-9,.\-]", "", texto)
    if not texto or texto.startswith("-") or "-" in texto:
        raise ErrorIngesta(f"{campo.capitalize()} vacio o negativo")

    coma = texto.rfind(",")
    punto = texto.rfind(".")
    if coma >= 0 and punto >= 0:
        decimal_sep = "," if coma > punto else "."
        millares = "." if decimal_sep == "," else ","
        texto = texto.replace(millares, "").replace(decimal_sep, ".")
    elif coma >= 0 or punto >= 0:
        separador = "," if coma >= 0 else "."
        partes = texto.split(separador)
        if len(partes) > 2:
            if all(len(parte) == 3 for parte in partes[1:]):
                texto = "".join(partes)
            else:
                raise ErrorIngesta(f"{campo.capitalize()} tiene separadores invalidos")
        elif len(partes) == 2 and len(partes[1]) == 3 and partes[0]:
            texto = "".join(partes)
        else:
            texto = ".".join(partes)

    try:
        numero = Decimal(texto).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as exc:
        raise ErrorIngesta(f"{campo.capitalize()} invalido") from exc
    if not numero.is_finite() or numero < 0 or numero > Decimal("99999999.99"):
        raise ErrorIngesta(f"{campo.capitalize()} fuera del rango permitido")
    return numero


def _producto_desde_fila(
    fila: Iterable[Any], columnas: dict[str, int], numero_fila: int
) -> dict[str, Any]:
    valores = list(fila)

    def obtener(campo: str) -> Any:
        indice = columnas.get(campo)
        return valores[indice] if indice is not None and indice < len(valores) else None

    try:
        codigo = _normalizar_texto(obtener("codigo")).upper()
        nombre = _normalizar_texto(obtener("nombre"))
        marca = _normalizar_texto(obtener("marca")) or None
        descripcion = _normalizar_texto(obtener("descripcion")) or None
    except ErrorIngesta as exc:
        raise ErrorIngesta(f"Fila {numero_fila}: {exc}") from exc
    if not codigo or not nombre:
        raise ErrorIngesta(f"Fila {numero_fila}: codigo y nombre son obligatorios")
    if len(codigo) > 50 or len(nombre) > 255:
        raise ErrorIngesta(f"Fila {numero_fila}: codigo o nombre excede el limite")

    if marca and len(marca) > 100:
        raise ErrorIngesta(f"Fila {numero_fila}: marca excede 100 caracteres")
    descuento = obtener("descuentos")
    try:
        precio = _leer_precio(obtener("precio"))
        descuentos = (
            _leer_precio(descuento, campo="descuento")
            if _texto(descuento)
            else Decimal("0.00")
        )
    except ErrorIngesta as exc:
        raise ErrorIngesta(f"Fila {numero_fila}: {exc}") from exc

    return {
        "codigo": codigo,
        "nombre": nombre,
        "marca": marca,
        "precio": precio,
        "descuentos": descuentos,
        "descripcion": descripcion,
    }


def _extraer_filas(
    filas: Iterable[Iterable[Any]],
    *,
    permitir_posicional: bool = False,
    desplazamiento: int = 0,
) -> tuple[list[dict[str, Any]], int, list[str]]:
    iterador = iter(filas)
    prefijo: list[list[Any]] = []
    indice_cabecera: int | None = None
    columnas: dict[str, int] | None = None
    for indice in range(100):
        try:
            fila = list(next(iterador))
        except StopIteration:
            break
        columnas = _mapear_cabecera(fila)
        if columnas is not None:
            indice_cabecera = indice
            break
        prefijo.append(fila)

    if columnas is None:
        if not permitir_posicional:
            raise CabeceraIngestaAusente(
                "No se encontro una cabecera reconocible de codigo, nombre y precio"
            )
        columnas = {campo: indice for indice, campo in enumerate(TXT_COLUMNAS_POSICIONALES)}
        indice_cabecera = -1

    filas_datos = iterador if indice_cabecera >= 0 else itertools.chain(prefijo, iterador)
    numero_inicio = indice_cabecera + 2 + desplazamiento
    if indice_cabecera == -1:
        numero_inicio = 1 + desplazamiento

    productos: list[dict[str, Any]] = []
    errores: list[str] = []
    cantidad_errores = 0
    cantidad_con_datos = 0
    cantidad_filas_leidas = 0
    for indice, fila in enumerate(filas_datos, start=numero_inicio):
        cantidad_filas_leidas += 1
        if cantidad_filas_leidas > MAX_FILAS_ARCHIVO + 101:
            raise ErrorIngesta(f"El archivo supera el maximo de {MAX_FILAS_ARCHIVO} filas")
        fila = list(fila)
        if not any(_texto(valor) for valor in fila):
            continue
        cantidad_con_datos += 1
        if cantidad_con_datos > MAX_FILAS_ARCHIVO:
            raise ErrorIngesta(f"El archivo supera el maximo de {MAX_FILAS_ARCHIVO} filas")
        try:
            if indice_cabecera == -1 and len(fila) > len(TXT_COLUMNAS_POSICIONALES):
                raise ErrorIngesta(
                    f"Fila {indice}: el orden posicional admite maximo "
                    f"{len(TXT_COLUMNAS_POSICIONALES)} columnas"
                )
            productos.append(_producto_desde_fila(fila, columnas, indice))
        except ErrorIngesta as exc:
            cantidad_errores += 1
            if len(errores) < MAX_ERRORES_DETALLE:
                errores.append(str(exc))

    if not productos and not cantidad_errores:
        raise ErrorIngesta("El archivo no contiene productos")
    if not productos:
        raise ArchivoSinFilasValidas(
            cantidad_errores,
            "No se encontro ninguna fila valida: " + "; ".join(errores),
        )
    return productos, cantidad_errores, errores


def _validar_paquete(contenido: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as archivo:
            entradas = archivo.infolist()
    except (zipfile.BadZipFile, OSError) as exc:
        raise ErrorIngesta("El archivo comprimido esta danado") from exc
    if len(entradas) > MAX_ENTRADAS_ZIP:
        raise ErrorIngesta("El archivo contiene demasiados elementos comprimidos")
    total = sum(entrada.file_size for entrada in entradas)
    if total > MAX_DESCOMPRIMIDO_ZIP:
        raise ErrorIngesta("El contenido descomprimido supera el limite permitido")
    for entrada in entradas:
        if (
            entrada.file_size > 1024 * 1024
            and entrada.file_size / max(entrada.compress_size, 1) > MAX_RAZON_COMPRESION
        ):
            raise ErrorIngesta("El archivo excede el limite de compresion seguro")


def _extraer_xlsx(contenido: bytes) -> tuple[list[dict[str, Any]], int, list[str]]:
    _validar_paquete(contenido)
    try:
        libro = load_workbook(
            io.BytesIO(contenido), read_only=True, data_only=True, keep_links=False
        )
    except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError, KeyError) as exc:
        raise ErrorIngesta("No se pudo leer el archivo Excel") from exc

    productos: list[dict[str, Any]] = []
    errores = 0
    detalles: list[str] = []
    try:
        for hoja in libro.worksheets:
            try:
                filas_hoja, errores_hoja, detalle_hoja = _extraer_filas(
                    hoja.iter_rows(values_only=True)
                )
            except CabeceraIngestaAusente:
                continue
            productos.extend(filas_hoja)
            errores += errores_hoja
            detalles.extend(detalle_hoja)
            if len(productos) + errores > MAX_FILAS_ARCHIVO:
                raise ErrorIngesta(
                    f"El archivo supera el maximo de {MAX_FILAS_ARCHIVO} filas"
                )
    finally:
        libro.close()
    if not productos:
        raise ErrorIngesta("El libro no contiene una tabla con codigo, nombre y precio")
    return productos, errores, detalles


def _extraer_docx(contenido: bytes) -> tuple[list[dict[str, Any]], int, list[str]]:
    _validar_paquete(contenido)
    try:
        documento = Document(io.BytesIO(contenido))
    except (PackageNotFoundError, zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ErrorIngesta("No se pudo leer el documento Word") from exc

    productos: list[dict[str, Any]] = []
    errores = 0
    detalles: list[str] = []
    for tabla in documento.tables:
        try:
            productos_tabla, errores_tabla, detalle_tabla = _extraer_filas(
                ([celda.text for celda in fila.cells] for fila in tabla.rows)
            )
        except CabeceraIngestaAusente:
            continue
        productos.extend(productos_tabla)
        errores += errores_tabla
        detalles.extend(detalle_tabla)
        if len(productos) + errores > MAX_FILAS_ARCHIVO:
            raise ErrorIngesta(
                f"El archivo supera el maximo de {MAX_FILAS_ARCHIVO} filas"
            )
    if not productos:
        raise ErrorIngesta("El documento no contiene una tabla con codigo, nombre y precio")
    return productos, errores, detalles


def _filas_pdf(texto: str) -> list[list[str]]:
    resultado: list[list[str]] = []
    for numero, linea in enumerate(io.StringIO(texto), start=1):
        if numero > MAX_FILAS_ARCHIVO + 101:
            raise ErrorIngesta(f"El archivo supera el maximo de {MAX_FILAS_ARCHIVO} filas")
        if not linea.strip():
            continue
        separador = _separador(linea)
        if separador:
            resultado.append(next(csv.reader([linea], delimiter=separador, skipinitialspace=True)))
        elif re.search(r"\s{2,}", linea.strip()):
            resultado.append(re.split(r"\s{2,}", linea.strip()))
        else:
            resultado.append(linea.split())
    return resultado


def _extraer_pdf(contenido: bytes) -> tuple[list[dict[str, Any]], int, list[str]]:
    try:
        lector = PdfReader(io.BytesIO(contenido), strict=True)
        if len(lector.pages) > MAX_PAGINAS_PDF:
            raise ErrorIngesta(f"El PDF supera el maximo de {MAX_PAGINAS_PDF} paginas")
        paginas = [pagina.extract_text() or "" for pagina in lector.pages]
    except ErrorIngesta:
        raise
    except (PdfReadError, OSError, ValueError, KeyError) as exc:
        raise ErrorIngesta("No se pudo leer el PDF") from exc

    texto = "\n".join(paginas)
    if not texto.strip():
        raise ErrorIngesta("El PDF no contiene texto legible; no se admite OCR")
    filas = _filas_pdf(texto)
    for indice, fila in enumerate(filas):
        columnas = _mapear_cabecera(fila)
        if columnas is None:
            continue
        encabezados = len(fila)
        filas_normalizadas = [fila]
        for datos in filas[indice + 1 :]:
            if len(datos) > encabezados and encabezados == 3 and set(columnas) >= {
                "codigo", "nombre", "precio"
            }:
                datos = [datos[0], " ".join(datos[1:-1]), datos[-1]]
            filas_normalizadas.append(datos)
        return _extraer_filas(filas_normalizadas)
    raise CabeceraIngestaAusente(
        "No se encontro una cabecera reconocible de codigo, nombre y precio"
    )


def _extraer_txt(contenido: bytes) -> tuple[list[dict[str, Any]], int, list[str]]:
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ErrorIngesta("El archivo TXT debe estar codificado en UTF-8") from exc
    filas = _filas_delimitadas(texto)
    primera_linea = texto.partition("\n")[0]
    tiene_separador_no_ambiguo = bool(_separador(primera_linea))
    tiene_cabecera = any(_mapear_cabecera(fila) for fila in filas[:100])
    if not tiene_cabecera and not tiene_separador_no_ambiguo and "," in texto:
        raise ErrorIngesta(
            "El TXT posicional debe separar columnas con punto y coma, barra vertical o tabulador"
        )
    return _extraer_filas(filas, permitir_posicional=True)


def _nombre_seguro(nombre: str) -> str:
    nombre = nombre.replace("\\", "/").rsplit("/", 1)[-1]
    nombre = "".join(caracter for caracter in nombre if caracter.isprintable())
    nombre = nombre.strip()
    if not nombre or len(nombre) > 255:
        raise ErrorIngesta("El nombre del archivo esta vacio o supera 255 caracteres")
    return nombre


def _ruta_original(nombre: str) -> tuple[Path, str]:
    directorio = get_settings().uploads_dir.resolve()
    estaticos = (Path(__file__).resolve().parents[1] / "static").resolve()
    if directorio == estaticos or estaticos in directorio.parents:
        raise RuntimeError("UPLOADS_DIR no puede estar dentro del directorio de estaticos")
    directorio.mkdir(parents=True, exist_ok=True)
    extension = Path(nombre).suffix.lower()
    nombre_archivo = f"{uuid.uuid4().hex}{extension}"
    ruta = directorio / nombre_archivo
    ruta_resuelta = ruta.resolve()
    if ruta_resuelta.parent != directorio:
        raise RuntimeError("La ruta configurada para cargas no es segura")
    return ruta, nombre_archivo


def _escribir_original(ruta: Path, contenido: bytes) -> None:
    try:
        descriptor = os.open(ruta, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(descriptor, "wb") as archivo:
            archivo.write(contenido)
    except OSError:
        ruta.unlink(missing_ok=True)
        raise


def _detalle(errores: list[str]) -> str | None:
    if not errores:
        return None
    detalle = "; ".join(errores[:MAX_ERRORES_DETALLE])
    return detalle[:MAX_DETALLE]


def _extraer(nombre: str, contenido: bytes) -> tuple[list[dict[str, Any]], int, list[str]]:
    extension = Path(nombre).suffix.lower()
    extractores = {
        ".pdf": _extraer_pdf,
        ".xlsx": _extraer_xlsx,
        ".docx": _extraer_docx,
        ".txt": _extraer_txt,
    }
    return extractores[extension](contenido)


def procesar_archivo(nombre: str, contenido: bytes, id_usuario: int) -> dict[str, Any]:
    """Procesa un archivo y devuelve un resultado aislado del resto del lote."""
    try:
        nombre = _nombre_seguro(nombre)
    except ErrorIngesta as exc:
        nombre_visible = str(nombre).replace("\\", "/").rsplit("/", 1)[-1]
        return {
            "id_archivo": None,
            "nombre": nombre_visible[:255] or "archivo",
            "tipo": Path(nombre_visible).suffix.lower().lstrip(".") or "desconocido",
            "tamano_bytes": len(contenido),
            "estado": "rechazado",
            "filas_ok": 0,
            "filas_error": 0,
            "detalle": str(exc),
        }

    extension = Path(nombre).suffix.lower()
    tipo = extension.lstrip(".") or "desconocido"
    if extension not in EXTENSIONES_PERMITIDAS:
        return {
            "id_archivo": None,
            "nombre": nombre,
            "tipo": tipo,
            "tamano_bytes": len(contenido),
            "estado": "rechazado",
            "filas_ok": 0,
            "filas_error": 0,
            "detalle": "Formato no permitido; use PDF, XLSX, DOCX o TXT",
        }
    if len(contenido) > MAX_ARCHIVO_BYTES:
        return {
            "id_archivo": None,
            "nombre": nombre,
            "tipo": tipo,
            "tamano_bytes": len(contenido),
            "estado": "rechazado",
            "filas_ok": 0,
            "filas_error": 0,
            "detalle": f"El archivo supera el limite de {MAX_ARCHIVO_MB} MB",
        }

    digest = hashlib.sha256(contenido).hexdigest()
    ruta, ruta_relativa = _ruta_original(nombre)
    _escribir_original(ruta, contenido)
    try:
        archivo = inventario_model.crear_archivo(
            nombre=nombre,
            tipo=tipo,
            ruta=ruta_relativa,
            hash_sha256=digest,
            tamano_bytes=len(contenido),
            id_usuario=id_usuario,
        )
    except (ErrorPostgres, BaseDeDatosNoDisponible, OSError):
        ruta.unlink(missing_ok=True)
        raise

    id_archivo = archivo["id_archivo"]
    inventario_model.actualizar_estado_archivo(id_archivo, estado="procesando")

    if not contenido:
        detalle = "El archivo esta vacio"
        resultado = inventario_model.actualizar_estado_archivo(
            id_archivo, estado="error", filas_error=1, detalle=detalle
        )
        return dict(resultado)

    try:
        productos, filas_error, errores = _extraer(nombre, contenido)
    except ArchivoSinFilasValidas as exc:
        resultado = inventario_model.actualizar_estado_archivo(
            id_archivo,
            estado="error",
            filas_error=exc.filas_error,
            detalle=str(exc)[:MAX_DETALLE],
        )
        return dict(resultado)
    except ErrorIngesta as exc:
        resultado = inventario_model.actualizar_estado_archivo(
            id_archivo, estado="error", filas_error=1, detalle=str(exc)[:MAX_DETALLE]
        )
        return dict(resultado)
    except (
        InvalidFileException,
        PackageNotFoundError,
        ParseError,
        PdfReadError,
        zipfile.BadZipFile,
        OSError,
        ValueError,
    ) as exc:
        logger.warning("No se pudo procesar archivo de ingesta %s: %s", nombre, exc)
        resultado = inventario_model.actualizar_estado_archivo(
            id_archivo,
            estado="error",
            filas_error=1,
            detalle="El archivo esta danado o no tiene un formato valido",
        )
        return dict(resultado)

    codigos: set[str] = set()
    unicos: list[dict[str, Any]] = []
    for producto in productos:
        codigo = producto["codigo"]
        if codigo in codigos:
            filas_error += 1
            if len(errores) < MAX_ERRORES_DETALLE:
                errores.append(f"Codigo duplicado en el archivo: {codigo}")
            continue
        codigos.add(codigo)
        unicos.append(producto)
    productos = unicos

    try:
        filas_ok = inventario_model.insertar_o_actualizar_productos(id_archivo, productos)
    except (ErrorPostgres, BaseDeDatosNoDisponible):
        inventario_model.actualizar_estado_archivo(
            id_archivo,
            estado="error",
            filas_error=max(filas_error, 1),
            detalle="No se pudieron guardar los productos del archivo",
        )
        raise

    estado = "parcial" if filas_error else "ok"
    resultado = inventario_model.actualizar_estado_archivo(
        id_archivo,
        estado=estado,
        filas_ok=filas_ok,
        filas_error=filas_error,
        detalle=_detalle(errores),
    )
    return dict(resultado)
